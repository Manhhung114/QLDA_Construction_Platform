import csv
import io
from collections import deque
from datetime import date, datetime
from typing import Any

from fastapi import APIRouter, Depends, File, HTTPException, Query, Request, UploadFile
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from . import models, models_extra
from .database import get_db
from .v2 import (
    APPROVABLE,
    REGISTRY,
    WORKFLOW,
    as_dict,
    audit,
    clean,
    current_user,
    model_for,
    permit,
    project_id_for,
)

router = APIRouter(prefix="/api", tags=["extended"])

IMPORTABLE = {
    "tasks", "schedule", "resources", "timesheets", "boq", "costs",
    "contracts", "payments", "procurement", "documents", "quality",
    "changes", "claims", "transmittals", "budget-versions",
}


def _pred_tokens(raw: str | None) -> list[str]:
    if not raw:
        return []
    text = raw.replace(";", ",").replace("|", ",")
    return [x.strip() for x in text.split(",") if x.strip()]


@router.get("/schedule/cpm")
def calculate_cpm(
    project_id: int,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    permit(db, user, "read", project_id, "schedule")
    activities = db.scalars(
        select(models.ScheduleActivity)
        .where(models.ScheduleActivity.project_id == project_id)
        .order_by(models.ScheduleActivity.id)
    ).all()
    if not activities:
        return {"project_id": project_id, "duration_days": 0, "activities": []}

    by_id = {str(a.id): a for a in activities}
    by_code = {str(a.activity_code).strip(): a for a in activities if a.activity_code}
    nodes: dict[int, dict[str, Any]] = {}
    for a in activities:
        if a.planned_start and a.planned_end:
            duration = max(1, (a.planned_end - a.planned_start).days + 1)
        else:
            duration = 1
        predecessors: list[int] = []
        for token in _pred_tokens(a.predecessor_ids):
            pred = by_id.get(token) or by_code.get(token)
            if pred and pred.id != a.id and pred.id not in predecessors:
                predecessors.append(pred.id)
        nodes[a.id] = {
            "activity": a,
            "duration": duration,
            "predecessors": predecessors,
            "successors": [],
        }

    for node_id, node in nodes.items():
        for pred_id in node["predecessors"]:
            if pred_id in nodes:
                nodes[pred_id]["successors"].append(node_id)

    indegree = {nid: len([p for p in n["predecessors"] if p in nodes]) for nid, n in nodes.items()}
    queue = deque([nid for nid, degree in indegree.items() if degree == 0])
    topo: list[int] = []
    while queue:
        nid = queue.popleft()
        topo.append(nid)
        for succ in nodes[nid]["successors"]:
            indegree[succ] -= 1
            if indegree[succ] == 0:
                queue.append(succ)
    if len(topo) != len(nodes):
        raise HTTPException(409, "Schedule dependencies contain a cycle")

    es: dict[int, int] = {}
    ef: dict[int, int] = {}
    for nid in topo:
        pred_finish = [ef[p] for p in nodes[nid]["predecessors"] if p in nodes]
        es[nid] = max(pred_finish, default=0)
        ef[nid] = es[nid] + nodes[nid]["duration"]
    project_duration = max(ef.values(), default=0)

    lf: dict[int, int] = {}
    ls: dict[int, int] = {}
    for nid in reversed(topo):
        succ_start = [ls[s] for s in nodes[nid]["successors"]]
        lf[nid] = min(succ_start, default=project_duration)
        ls[nid] = lf[nid] - nodes[nid]["duration"]

    rows = []
    for nid in topo:
        a = nodes[nid]["activity"]
        total_float = ls[nid] - es[nid]
        rows.append({
            **as_dict(a),
            "duration_days": nodes[nid]["duration"],
            "early_start_day": es[nid],
            "early_finish_day": ef[nid],
            "late_start_day": ls[nid],
            "late_finish_day": lf[nid],
            "total_float_days": total_float,
            "computed_critical": total_float == 0,
        })
    return {"project_id": project_id, "duration_days": project_duration, "activities": rows}


@router.get("/calendar")
def project_calendar(
    project_id: int,
    start: date | None = None,
    end: date | None = None,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    permit(db, user, "read", project_id, "calendar")
    events: list[dict[str, Any]] = []

    def inside(d: date | None) -> bool:
        if d is None:
            return False
        return (start is None or d >= start) and (end is None or d <= end)

    for t in db.scalars(select(models.Task).where(models.Task.project_id == project_id)).all():
        if inside(t.start_date):
            events.append({"date": t.start_date.isoformat(), "type": "task_start", "id": t.id, "title": f"Bắt đầu: {t.title}", "status": t.status})
        if inside(t.due_date):
            events.append({"date": t.due_date.isoformat(), "type": "task_due", "id": t.id, "title": f"Deadline: {t.title}", "status": t.status})
    for a in db.scalars(select(models.ScheduleActivity).where(models.ScheduleActivity.project_id == project_id)).all():
        if a.milestone and inside(a.planned_end or a.planned_start):
            d = a.planned_end or a.planned_start
            events.append({"date": d.isoformat(), "type": "milestone", "id": a.id, "title": a.name, "critical": a.critical})
    for q in db.scalars(select(models.QualityItem).where(models.QualityItem.project_id == project_id)).all():
        if inside(q.due_date):
            events.append({"date": q.due_date.isoformat(), "type": q.item_type.lower(), "id": q.id, "title": f"{q.item_type}: {q.title}", "status": q.status})
    for c in db.scalars(select(models.Contract).where(models.Contract.project_id == project_id)).all():
        if inside(c.end_date):
            events.append({"date": c.end_date.isoformat(), "type": "contract_end", "id": c.id, "title": f"Hết hạn HĐ: {c.contract_no}", "status": c.status})
    for p in db.scalars(select(models.ProcurementOrder).where(models.ProcurementOrder.project_id == project_id)).all():
        if inside(p.expected_date):
            events.append({"date": p.expected_date.isoformat(), "type": "procurement", "id": p.id, "title": f"Giao hàng: {p.po_no}", "status": p.status})
    events.sort(key=lambda x: (x["date"], x["type"], x["id"]))
    return events


def _sheet_rows(content: bytes, filename: str) -> list[dict[str, Any]]:
    lower = filename.lower()
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig")
        return [dict(r) for r in csv.DictReader(io.StringIO(text))]
    if lower.endswith(".xlsx"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            headers = [str(x).strip() if x is not None else "" for x in next(rows)]
        except StopIteration:
            return []
        result = []
        for values in rows:
            item = {headers[i]: values[i] for i in range(min(len(headers), len(values))) if headers[i]}
            if any(v not in (None, "") for v in item.values()):
                result.append(item)
        return result
    raise HTTPException(415, "Import supports CSV or XLSX")


@router.post("/import/{module}")
async def import_records(
    module: str,
    request: Request,
    file: UploadFile = File(...),
    project_id: int | None = Query(default=None),
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    if module not in IMPORTABLE:
        raise HTTPException(400, f"Import not enabled for {module}")
    model = model_for(module)
    filename = file.filename or "import.csv"
    content = await file.read()
    max_mb = int(__import__("os").getenv("MAX_IMPORT_MB", "20"))
    if len(content) > max_mb * 1024 * 1024:
        raise HTTPException(413, f"Import file exceeds {max_mb} MB")
    raw_rows = _sheet_rows(content, filename)
    created = 0
    errors: list[dict[str, Any]] = []
    for index, raw in enumerate(raw_rows, start=2):
        payload = {str(k).strip(): v for k, v in raw.items() if str(k).strip()}
        if project_id is not None and hasattr(model, "project_id"):
            payload["project_id"] = project_id
        pid = project_id_for(module, payload=payload)
        try:
            permit(db, user, "write", pid, module)
            data = clean(model, payload)
            if module == "contracts" and data.get("start_date") and data.get("duration_days") and not data.get("end_date"):
                from datetime import timedelta
                data["end_date"] = data["start_date"] + timedelta(days=int(data["duration_days"]))
            row = model(**data)
            db.add(row)
            db.flush()
            audit(db, user, module, row.id, "import_create", None, as_dict(row), request)
            created += 1
        except Exception as exc:
            db.rollback()
            errors.append({"row": index, "error": exc.__class__.__name__, "detail": str(exc)[:300]})
    return {"module": module, "created": created, "errors": errors, "total_rows": len(raw_rows)}


@router.post("/automation/rules/run")
def run_saved_rules(
    request: Request,
    project_id: int | None = None,
    db: Session = Depends(get_db),
    user=Depends(current_user),
):
    query = select(models.WorkflowRule).where(models.WorkflowRule.enabled.is_(True))
    if project_id is not None:
        permit(db, user, "write", project_id, "automation")
        query = query.where((models.WorkflowRule.project_id == project_id) | (models.WorkflowRule.project_id.is_(None)))
    rules = db.scalars(query.order_by(models.WorkflowRule.id)).all()
    changed = 0
    skipped = 0
    details: list[dict[str, Any]] = []
    for rule in rules:
        model = REGISTRY.get(rule.module)
        if not model or not hasattr(model, "status"):
            skipped += 1
            continue
        stmt = select(model).where(model.status == rule.trigger_status)
        if rule.project_id is not None and hasattr(model, "project_id"):
            stmt = stmt.where(model.project_id == rule.project_id)
        elif project_id is not None and hasattr(model, "project_id"):
            stmt = stmt.where(model.project_id == project_id)
        rows = db.scalars(stmt.limit(500)).all()
        for row in rows:
            pid = project_id_for(rule.module, row, db=db)
            try:
                permit(db, user, "write", pid, rule.module)
            except HTTPException:
                skipped += 1
                continue
            if rule.module in APPROVABLE and rule.action_status not in WORKFLOW.get(rule.module, {}).get(row.status, []):
                skipped += 1
                continue
            before = as_dict(row)
            row.status = rule.action_status
            db.add(models_extra.ApprovalAction(
                project_id=pid,
                module=rule.module,
                entity_id=row.id,
                from_status=rule.trigger_status,
                to_status=rule.action_status,
                action="automation",
                comment=f"Workflow rule #{rule.id}",
                actor_user_id=user.id,
                actor_name=user.full_name or user.username,
            ))
            db.add(models.Notification(
                user_id=None,
                project_id=pid,
                title=f"Workflow tự động: {rule.module} #{row.id}",
                message=f"{rule.trigger_status} → {rule.action_status} bởi rule #{rule.id}",
            ))
            db.commit()
            audit(db, user, rule.module, row.id, "automation_transition", before, as_dict(row), request)
            changed += 1
            details.append({"rule_id": rule.id, "module": rule.module, "entity_id": row.id, "to_status": rule.action_status})
    return {"rules": len(rules), "changed": changed, "skipped": skipped, "details": details[:200], "run_at": datetime.utcnow().isoformat()}
